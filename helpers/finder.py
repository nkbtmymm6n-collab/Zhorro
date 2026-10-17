import sqlite3
from pathlib import Path
from docx import Document
from .appender import appender
import pymupdf

def finder(word, model, limit, selected):
    db=sqlite3.connect("zhorro.db")

    list_of_founds=[]
    list_of_results=[]
    files=selected


    for file in files:
        path=Path(file)
        filename=path.name
        founds={}
        founds["path"]=path
        founds["filename"]=filename
        founds["contexts"]=[]
        list_of_founds.append(founds)

    for file in list_of_founds:
        if Path(file["path"]).suffix==".txt":
            with open (file["path"], "r", encoding="utf-8") as f:
                text=f.read()
                positions=[{"start": 0, "end": len(text)}]
            None, list_of_results.extend(appender(text, file, word, positions, model, True, limit))

        elif Path(file["path"]).suffix==".pdf":
            doc=pymupdf.open(file["path"])
            positions=[]
            number=0
            counter=""
            text=""
            for page in doc:
                text+=page.get_text()
                positions.append({"start": len(counter), "end": len(text), "page": number})
                counter=text
                number+=1
        
            None, list_of_results.extend(appender(text, file, word, positions, model, True, limit))

        elif Path(file["path"]).suffix==".docx":
            doc=Document(file["path"])
            paragraphs=doc.paragraphs
            positions=[]
            number=0
            counter=""
            text=""

            for paragraph in paragraphs:
                text+=paragraph.text
                positions.append({"start": len(counter), "end": len(text), "paragraph": number})
                counter=text
                number+=1
                
            None, list_of_results.extend(appender(text, file, word, positions, model, True, limit))                

    all_files=selected
    for file in all_files:
        path=Path(file)
        filename=path.name
        results={}
        results["path"]=path
        results["filename"]=filename

        if results["path"].suffix==".txt":
            with open (results["path"], "r", encoding="utf-8") as f:
                text=f.read()
                list_of_results.extend(appender(text, results, word, None, model, False, limit))

        elif results["path"].suffix==".pdf":
            doc=pymupdf.open(results["path"])
            positions=[]
            counter=""
            text=""               
            number=0

            for page in doc:
                text+=page.get_text()
                positions.append({"start": len(counter), "end": len(text), "page": number})
                counter=text
                number+=1
                
            list_of_results.extend(appender(text, results, word, positions, model, False, limit))

        elif results["path"].suffix==".docx":
            doc=Document(results["path"])
            paragraphs=doc.paragraphs
            positions=[]
            counter=""
            text=""
            number=0

            for paragraph in paragraphs:
                text+=paragraph.text
                positions.append({"start": len(counter), "end": len(text), "paragraph": number})
                number+=1

            list_of_results.extend(appender(text, results, word, positions, model, False, limit))

    for file in list_of_results:
        file.sort(key=lambda x: x["similarity"], reverse=True)
        for result in file:
            result["path"]=str(result["path"])

    for file in list_of_founds:
        file["path"]=str(file["path"])

    db.close()
    return list_of_founds, list_of_results

